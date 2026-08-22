import io
import logging
from datetime import datetime, timezone
from fastapi import APIRouter, Depends, HTTPException, status, Body
from fastapi.responses import StreamingResponse
from app.auth import get_current_user
from app.db import get_document_by_id, list_user_documents, update_document, delete_document_by_id

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api", tags=["documents"])


@router.get("/documents")
async def list_documents(current_user: dict = Depends(get_current_user)):
    """
    Returns all documents owned by the authenticated user.
    """
    try:
        results = list_user_documents(current_user["uid"])
        for d in results:
            for ts_field in ("upload_date", "updated_at"):
                if d.get(ts_field) and hasattr(d[ts_field], "isoformat"):
                    d[ts_field] = d[ts_field].isoformat()
        return {"documents": results}
    except Exception as e:
        logger.exception(f"Error listing documents: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to list documents: {str(e)}"
        )


@router.get("/documents/{document_id}")
async def get_document(document_id: str, current_user: dict = Depends(get_current_user)):
    """
    Returns a single document by ID, if owned by the authenticated user.
    """
    d = get_document_by_id(document_id)

    if not d:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found.")

    if d.get("user_id") != current_user["uid"]:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied.")

    for ts_field in ("upload_date", "updated_at"):
        if d.get(ts_field) and hasattr(d[ts_field], "isoformat"):
            d[ts_field] = d[ts_field].isoformat()

    return d


@router.put("/documents/{document_id}/text")
async def update_document_text(
    document_id: str,
    body: dict = Body(...),
    current_user: dict = Depends(get_current_user)
):
    """
    Updates the recognized_text for a document after user edits.
    """
    d = get_document_by_id(document_id)

    if not d:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found.")

    if d.get("user_id") != current_user["uid"]:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied.")

    new_text = body.get("recognized_text", "")
    update_document(document_id, {
        "recognized_text": new_text,
        "updated_at": datetime.now(timezone.utc)
    })

    return {"document_id": document_id, "status": "updated"}


@router.delete("/documents/{document_id}")
async def delete_document(document_id: str, current_user: dict = Depends(get_current_user)):
    """
    Deletes a document record.
    """
    d = get_document_by_id(document_id)

    if not d:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found.")

    if d.get("user_id") != current_user["uid"]:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied.")

    delete_document_by_id(document_id)
    return {"document_id": document_id, "status": "deleted"}


@router.get("/documents/{document_id}/export")
async def export_document(
    document_id: str,
    format: str = "txt",
    current_user: dict = Depends(get_current_user)
):
    """
    Exports the document's recognized text in TXT, DOCX, or PDF format.
    """
    d = get_document_by_id(document_id)

    if not d:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found.")

    if d.get("user_id") != current_user["uid"]:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied.")

    text = d.get("recognized_text", "")
    doc_title = f"Scriptly_Document_{document_id[:8]}"

    if format == "txt":
        return StreamingResponse(
            io.BytesIO(text.encode("utf-8")),
            media_type="text/plain",
            headers={"Content-Disposition": f'attachment; filename="{doc_title}.txt"'}
        )

    elif format == "docx":
        try:
            from docx import Document
            from docx.shared import Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            document = Document()
            title = document.add_heading("Scriptly — Transcription Result", level=1)
            title.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
            meta_para = document.add_paragraph()
            meta_para.add_run(f"Document ID: {document_id}").italic = True
            meta_para.paragraph_format.space_after = Pt(12)

            document.add_paragraph("─" * 50)

            for para_text in text.split("\n"):
                p = document.add_paragraph(para_text)
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)

            buf = io.BytesIO()
            document.save(buf)
            buf.seek(0)

            return StreamingResponse(
                buf,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": f'attachment; filename="{doc_title}.docx"'}
            )
        except Exception as e:
            logger.exception(f"DOCX generation failed: {e}")
            raise HTTPException(status_code=500, detail=f"DOCX export failed: {str(e)}")

    elif format == "pdf":
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import inch
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

            buf = io.BytesIO()
            doc = SimpleDocTemplate(buf, pagesize=A4,
                                    leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                                    topMargin=0.8 * inch, bottomMargin=0.8 * inch)
            
            styles = getSampleStyleSheet()
            title_style = styles["Title"]
            body_style = ParagraphStyle(
                "BodyCustom",
                parent=styles["Normal"],
                fontSize=11,
                leading=16,
                spaceAfter=6
            )
            meta_style = ParagraphStyle(
                "Meta",
                parent=styles["Normal"],
                fontSize=9,
                textColor="#5B5F6B",
                spaceAfter=12
            )

            elements = []
            elements.append(Paragraph("Scriptly — Transcription Result", title_style))
            elements.append(Paragraph(f"Document ID: {document_id}", meta_style))
            elements.append(Spacer(1, 12))

            import html
            for line in text.split("\n"):
                safe_line = html.escape(line) if line.strip() else "&nbsp;"
                elements.append(Paragraph(safe_line, body_style))

            doc.build(elements)
            buf.seek(0)

            return StreamingResponse(
                buf,
                media_type="application/pdf",
                headers={"Content-Disposition": f'attachment; filename="{doc_title}.pdf"'}
            )
        except Exception as e:
            logger.exception(f"PDF generation failed: {e}")
            raise HTTPException(status_code=500, detail=f"PDF export failed: {str(e)}")

    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported export format '{format}'. Supported: txt, docx, pdf"
        )
